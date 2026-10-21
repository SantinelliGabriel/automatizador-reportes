import tempfile
import sqlite3
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import database
import main
import reportes
from validaciones import validar_ventas



class TestValidacionesVentas(unittest.TestCase):

    def test_rechaza_precio_no_numerico(self):
        ventas = self.ventas.copy()
        ventas["precio_unitario"] = ventas["precio_unitario"].astype(object)
        ventas.loc[0, "precio_unitario"] = "abc"

        with self.assertRaisesRegex(
            ValueError,
            "no numéricos",
        ):
            validar_ventas(ventas)

    def test_rechaza_unidades_decimales(self):
        ventas = self.ventas.copy()
        ventas["unidades"] = ventas["unidades"].astype(float)
        ventas.loc[0, "unidades"] = 2.5

        with self.assertRaisesRegex(
            ValueError,
            "números enteros",
        ):
            validar_ventas(ventas)

    def test_acepta_precio_cero(self):
        self.ventas.loc[0, "precio_unitario"] = 0

        resultado = validar_ventas(self.ventas)

        self.assertEqual(resultado.loc[0, "precio_unitario"], 0)

    def test_redondea_precio_a_centavos(self):
        ventas = self.ventas.copy()
        ventas["precio_unitario"] = ventas["precio_unitario"].astype(float)
        ventas.loc[0, "precio_unitario"] = 1.005

        resultado = validar_ventas(ventas)

        self.assertEqual(resultado.loc[0, "precio_unitario_centavos"], 101)
        self.assertEqual(resultado.loc[0, "precio_unitario"], 1.01)

    def setUp(self):
        # Datos correctos para reutilizar en las pruebas
        self.ventas = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

    def test_acepta_datos_correctos(self):
        resultado = validar_ventas(self.ventas.copy())
        self.assertEqual(len(resultado), 2)

    def test_rechaza_precio_negativo(self):
        self.ventas.loc[0, "precio_unitario"] = -12000

        with self.assertRaisesRegex(ValueError, "precios no pueden ser negativos"):
            validar_ventas(self.ventas)

    def test_rechaza_unidades_cero(self):
        self.ventas.loc[0, "unidades"] = 0

        with self.assertRaisesRegex(ValueError, "unidades deben ser mayores que cero"):
            validar_ventas(self.ventas)

    def test_rechaza_columna_faltante(self):
        ventas = self.ventas.drop(columns=["producto"])

        with self.assertRaisesRegex(ValueError, "Faltan estas columnas"):
            validar_ventas(ventas)

    def test_rechaza_celdas_vacias(self):
        self.ventas.loc[0, "producto"] = None

        with self.assertRaisesRegex(ValueError, "celdas vacías"):
            validar_ventas(self.ventas)

    def test_elimina_espacios_del_nombre_producto(self):
        ventas = self.ventas.copy()
        ventas.loc[0, "producto"] = "  Teclado  "

        resultado = validar_ventas(ventas)

        self.assertEqual(resultado.loc[0, "producto"], "Teclado")

    def test_no_modifica_el_dataframe_original(self):
        ventas = self.ventas.copy()
        ventas.loc[0, "producto"] = "  Teclado  "

        validar_ventas(ventas)

        self.assertEqual(ventas.loc[0, "producto"], "  Teclado  ")
        self.assertNotIn("precio_unitario_centavos", ventas.columns)

    def test_rechaza_precio_infinito(self):
        ventas = self.ventas.copy()
        ventas["precio_unitario"] = ventas["precio_unitario"].astype(float)
        ventas.loc[0, "precio_unitario"] = float("inf")

        with self.assertRaisesRegex(ValueError, "números no válidos"):
            validar_ventas(ventas)

    def test_rechaza_unidades_nan(self):
        ventas = self.ventas.copy()
        ventas.loc[0, "unidades"] = float("nan")

        with self.assertRaisesRegex(ValueError, "celdas vacías"):
            validar_ventas(ventas)

    def test_rechaza_importe_que_desborda(self):
        ventas = self.ventas.copy()
        ventas["unidades"] = ventas["unidades"].astype(float)
        ventas["precio_unitario"] = ventas["precio_unitario"].astype(float)

        ventas.loc[0, "unidades"] = 1e9
        ventas.loc[0, "precio_unitario"] = 1e10

        with self.assertRaisesRegex(
            ValueError,
            "importe calculado por producto no es válido",
        ):
            validar_ventas(ventas)

    def test_rechaza_unidades_fuera_del_rango_de_sqlite(self):
        ventas = self.ventas.copy()
        ventas["unidades"] = ventas["unidades"].astype(object)
        ventas.loc[0, "unidades"] = 2**63

        with self.assertRaisesRegex(
            ValueError,
            "unidades están fuera del rango permitido",
        ):
            validar_ventas(ventas)

class TestBaseDeDatos(unittest.TestCase):

    def setUp(self):
        """Prepara una base de datos temporal para cada prueba."""
        self.carpeta_temporal = tempfile.TemporaryDirectory()
        self.archivo_bd = Path(self.carpeta_temporal.name) / "pruebas.db"

        self.parche_bd = patch.object(database, "ARCHIVO_BD", self.archivo_bd)
        self.parche_bd.start()

        self.addCleanup(self.parche_bd.stop)
        self.addCleanup(self.carpeta_temporal.cleanup)

        database.crear_tabla()

    def test_crea_tabla_ventas(self):
        """Comprueba que se crea la tabla de ventas."""
        conexion = database.conectar_bd()

        try:
            resultado = conexion.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'ventas'
                """
            ).fetchone()
        finally:
            conexion.close()

        self.assertIsNotNone(resultado)


    def test_guarda_ventas_en_la_base_de_datos(self):
        """Comprueba que las ventas se guardan correctamente."""
        ventas = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

        database.guardar_ventas(ventas)
        resultado = database.consultar_ventas()

        self.assertEqual(len(resultado), 2)
        self.assertEqual(resultado["producto"].tolist(), ["Teclado", "Mouse"])
        self.assertEqual(resultado["total"].tolist(), [60000, 60000])


    def test_guardar_dos_veces_reemplaza_las_ventas(self):
        """Comprueba que una nueva importación sustituya la anterior."""
        ventas = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

        database.guardar_ventas(ventas)
        nuevas_ventas = pd.DataFrame({
            "producto": ["Monitor"],
            "unidades": [2],
            "precio_unitario": [150000],
        })
        database.guardar_ventas(nuevas_ventas)

        resultado = database.consultar_ventas()

        self.assertEqual(len(resultado), 1)
        self.assertEqual(resultado["producto"].tolist(), ["Monitor"])


    def test_base_de_datos_inicialmente_vacia(self):
        """Comprueba que la tabla no tenga ventas al crearse."""
        resultado = database.consultar_ventas()

        self.assertTrue(resultado.empty)
        self.assertEqual(
            list(resultado.columns),
            [
                "producto", "unidades", "precio_unitario",
                "total_centavos", "total",
            ],
        )


    def test_error_al_guardar_conserva_las_ventas_anteriores(self):
        """Comprueba que un error deshaga la operación incompleta."""
        ventas_originales = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

        ventas_nuevas = pd.DataFrame({
            "producto": ["Monitor"],
            "unidades": [2],
            "precio_unitario": [150000],
        })

        database.guardar_ventas(ventas_originales)

        with patch.object(
            database,
            "conectar_bd",
            side_effect=RuntimeError("Error simulado"),
        ):
            with self.assertRaisesRegex(RuntimeError, "Error simulado"):
                database.guardar_ventas(ventas_nuevas)

        resultado = database.consultar_ventas()

        self.assertEqual(
            resultado["producto"].tolist(),
            ["Teclado", "Mouse"],
        )
        self.assertEqual(len(resultado), 2)


    def test_rollback_restaura_ventas_tras_error_de_insercion(self):
        """Comprueba que rollback restaure los datos anteriores."""
        ventas_originales = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

        ventas_nuevas = pd.DataFrame({
            "producto": ["Monitor"],
            "unidades": [2],
            "precio_unitario": [150000],
        })

        database.guardar_ventas(ventas_originales)

        conexion_original = database.conectar_bd

        def conectar_con_error():
            conexion = conexion_original()

            conexion.execute(
                """
                CREATE TRIGGER provocar_error
                BEFORE INSERT ON ventas
                BEGIN
                    SELECT RAISE(ABORT, 'Error de inserción simulado');
                END
                """
            )

            conexion.commit()
            return conexion

        with patch.object(
            database,
            "conectar_bd",
            side_effect=conectar_con_error,
        ):
            with self.assertRaisesRegex(
                Exception,
                "Error de inserción simulado",
            ):
                database.guardar_ventas(ventas_nuevas)

        resultado = database.consultar_ventas()

        self.assertEqual(
            resultado["producto"].tolist(),
            ["Teclado", "Mouse"],
        )
        self.assertEqual(len(resultado), 2)

    def test_rollback_si_falla_la_publicacion_del_reporte(self):
        """Una falla al publicar el reporte conserva las ventas anteriores."""
        originales = pd.DataFrame({
            "producto": ["Teclado"], "unidades": [5], "precio_unitario": [12000],
        })
        nuevas = pd.DataFrame({
            "producto": ["Monitor"], "unidades": [2], "precio_unitario": [150000],
        })
        database.guardar_ventas(originales)

        def fallar_publicacion():
            raise RuntimeError("falló publicación")

        with self.assertRaisesRegex(RuntimeError, "falló publicación"):
            database.guardar_ventas(
                nuevas, antes_de_commit=lambda _ventas: fallar_publicacion()
            )

        resultado = database.consultar_ventas()
        self.assertEqual(resultado["producto"].tolist(), ["Teclado"])

    def test_calcula_subtotales_correctamente(self):
        """Comprueba que SQL calcule los importes de cada venta."""
        ventas = pd.DataFrame({
            "producto": ["Teclado", "Mouse"],
            "unidades": [5, 8],
            "precio_unitario": [12000, 7500],
        })

        database.guardar_ventas(ventas)
        resultado = database.consultar_ventas()

        self.assertEqual(resultado["total"].tolist(), [60000, 60000])
        self.assertEqual(resultado["total"].sum(), 120000)

    def test_migra_base_existente_a_centavos(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "legacy.db"
            conexion = sqlite3.connect(ruta)
            conexion.execute(
                """CREATE TABLE ventas (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    producto TEXT NOT NULL,
                    unidades INTEGER NOT NULL,
                    precio_unitario REAL NOT NULL
                )"""
            )
            conexion.execute(
                "INSERT INTO ventas (producto, unidades, precio_unitario) VALUES (?, ?, ?)",
                ("Adaptador", 2, 10.15),
            )
            conexion.commit()
            conexion.close()

            database.crear_tabla(ruta)
            resultado = database.consultar_ventas(ruta)

            self.assertEqual(resultado.loc[0, "precio_unitario"], 10.15)
            self.assertEqual(resultado.loc[0, "total_centavos"], 2030)
            self.assertEqual(resultado.loc[0, "total"], 20.3)

class TestReportes(unittest.TestCase):

    def test_aplica_formato_al_excel(self):
        """Comprueba que el reporte recibe el formato esperado."""

        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / "reporte.xlsx"

            ventas = pd.DataFrame({
                "producto": ["Teclado", "Mouse"],
                "unidades": [2, 3],
                "precio_unitario": [10000.0, 5000.0],
                "total": [20000.0, 15000.0],
            })

            ventas.to_excel(archivo, index=False)
            reportes.aplicar_formato_excel(archivo)

            libro = load_workbook(archivo)

            try:
                hoja = libro.active

                self.assertEqual(hoja["A1"].value, "producto")
                self.assertEqual(hoja.freeze_panes, "A2")
                self.assertEqual(hoja["A1"].fill.fgColor.rgb, "001F4E78")
                self.assertEqual(hoja["C2"].number_format, '"$"#,##0.00')
                self.assertEqual(hoja.auto_filter.ref, "A1:D2")
            finally:
                libro.close()

    def test_formatea_correctamente_la_fila_total(self):
        """Comprueba el formato de la última fila del reporte."""

        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / "reporte.xlsx"

            ventas = pd.DataFrame({
                "producto": ["Teclado", "Mouse", "TOTAL GENERAL"],
                "unidades": [2, 3, 5],
                "precio_unitario": [10000.0, 5000.0, 0],
                "total": [20000.0, 15000.0, 35000.0],
            })

            ventas.to_excel(archivo, index=False)
            reportes.aplicar_formato_excel(archivo)

            libro = load_workbook(archivo)

            try:
                hoja = libro.active
                ultima_fila = hoja.max_row

                self.assertEqual(
                    hoja.cell(ultima_fila, 1).value,
                    "TOTAL GENERAL",
                )
                self.assertTrue(
                    hoja.cell(ultima_fila, 1).font.bold
                )
                self.assertEqual(
                    hoja.cell(ultima_fila, 1).fill.fgColor.rgb,
                    "00D9EAF7",
                )
                self.assertEqual(
                    hoja.auto_filter.ref,
                    f"A1:D{ultima_fila - 1}",
                )
            finally:
                libro.close()

    def test_ajusta_ancho_de_columnas(self):
        """Comprueba que las columnas se ajusten al contenido."""

        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / "reporte.xlsx"

            ventas = pd.DataFrame({
                "producto": ["Teclado mecánico profesional"],
                "unidades": [2],
                "precio_unitario": [10000.0],
                "total": [20000.0],
            })

            ventas.to_excel(archivo, index=False)
            reportes.aplicar_formato_excel(archivo)

            libro = load_workbook(archivo)

            try:
                hoja = libro.active

                ancho_esperado = len("Teclado mecánico profesional") + 4

                self.assertGreaterEqual(
                    hoja.column_dimensions["A"].width,
                    ancho_esperado,
                )
            finally:
                libro.close()

    def test_formatea_encabezados_correctamente(self):
        """Comprueba el estilo y la altura de los encabezados."""

        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / "reporte.xlsx"

            ventas = pd.DataFrame({
                "producto": ["Teclado"],
                "unidades": [2],
                "precio_unitario": [10000.0],
                "total": [20000.0],
            })

            ventas.to_excel(archivo, index=False)
            reportes.aplicar_formato_excel(archivo)

            libro = load_workbook(archivo)

            try:
                hoja = libro.active
                encabezado = hoja["A1"]

                self.assertTrue(encabezado.font.bold)
                self.assertEqual(encabezado.font.color.type, "rgb")
                self.assertEqual(encabezado.font.color.rgb, "00FFFFFF")
                self.assertEqual(encabezado.fill.fgColor.rgb, "001F4E78")
                self.assertEqual(hoja.row_dimensions[1].height, 24)
            finally:
                libro.close()

    def test_generar_reporte_completo(self):
        """Comprueba que se genere el reporte con el total correcto."""

        with tempfile.TemporaryDirectory() as carpeta:
            carpeta_temporal = Path(carpeta)
            archivo_entrada = carpeta_temporal / "entrada.xlsx"
            archivo_salida = carpeta_temporal / "reporte.xlsx"
            archivo_bd = carpeta_temporal / "pruebas.db"

            ventas_entrada = pd.DataFrame({
                "producto": ["Teclado", "Mouse"],
                "unidades": [2, 3],
                "precio_unitario": [10000.0, 5000.0],
            })
            ventas_entrada.to_excel(archivo_entrada, index=False)

            with (
                patch.object(main, "ARCHIVO_ENTRADA", archivo_entrada),
                patch.object(main, "ARCHIVO_SALIDA", archivo_salida),
                patch.object(database, "ARCHIVO_BD", archivo_bd),
            ):
                main.generar_reporte()

            self.assertTrue(archivo_salida.exists())

            reporte = pd.read_excel(archivo_salida)

            self.assertEqual(len(reporte), 3)
            self.assertEqual(reporte.iloc[-1]["producto"], "TOTAL GENERAL")
            self.assertEqual(reporte.iloc[-1]["unidades"], 5)
            self.assertEqual(reporte.iloc[-1]["total"], 35000.0)


class TestConsola(unittest.TestCase):

    def test_main_pasa_las_rutas_configuradas(self):
        with patch.object(main, "generar_reporte") as generar:
            codigo = main.main([
                "--input", "entrada.xlsx",
                "--output", "salida.xlsx",
                "--database", "ventas.db",
            ])

        self.assertEqual(codigo, 0)
        generar.assert_called_once_with(
            Path("entrada.xlsx"), Path("salida.xlsx"), Path("ventas.db")
        )

    def test_main_devuelve_error_si_falla_el_proceso(self):
        with patch.object(main, "generar_reporte", side_effect=RuntimeError("fallo")):
            codigo = main.main([])

        self.assertEqual(codigo, 1)
if __name__ == "__main__":
    unittest.main(verbosity=2)
