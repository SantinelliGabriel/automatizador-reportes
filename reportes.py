
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd


def construir_reporte(ventas):
    """Devuelve las ventas con su fila de total y el importe general."""
    total_centavos = sum(int(valor) for valor in ventas["total_centavos"])
    total_ventas = total_centavos / 100
    ventas_reporte = ventas[
        ["producto", "unidades", "precio_unitario", "total"]
    ].copy()
    fila_total = {
        "producto": "TOTAL GENERAL",
        "unidades": sum(int(valor) for valor in ventas["unidades"]),
        "precio_unitario": None,
        "total": total_ventas,
    }
    return (
        pd.concat([ventas_reporte, pd.DataFrame([fila_total])], ignore_index=True),
        total_ventas,
    )


def aplicar_formato_excel(ruta_archivo):
    """Aplica formato profesional al reporte generado."""

    libro = load_workbook(ruta_archivo)
    hoja = libro.active

    azul_oscuro = "1F4E78"
    azul_claro = "D9EAF7"
    blanco = "FFFFFF"

    for celda in hoja[1]:
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor=azul_oscuro,
        )
        celda.font = Font(bold=True, color=blanco, size=11)
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(
            len(str(celda.value or ""))
            for celda in columna
        )
        hoja.column_dimensions[letra].width = ancho + 4

    for fila in range(2, hoja.max_row + 1):
        hoja.cell(fila, 3).number_format = '"$"#,##0.00'
        hoja.cell(fila, 4).number_format = '"$"#,##0.00'

        for columna in [2, 3, 4]:
            hoja.cell(fila, columna).alignment = Alignment(
                horizontal="right"
            )

    for celda in hoja[hoja.max_row]:
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor=azul_claro,
        )
        celda.font = Font(bold=True, color=azul_oscuro)
        celda.border = Border(
            top=Side(style="thin", color=azul_oscuro)
        )

    hoja.freeze_panes = "A2"

    # Los filtros abarcan los encabezados y los productos,
    # sin incluir la fila TOTAL GENERAL.
    hoja.auto_filter.ref = f"A1:D{hoja.max_row - 1}"
    hoja.row_dimensions[1].height = 24

    libro.save(ruta_archivo)
    libro.close()
