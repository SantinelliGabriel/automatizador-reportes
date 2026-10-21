"""Construye el libro consolidado de comprobantes y observaciones."""

import json

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd


def construir_hojas(comprobantes):
    if comprobantes.empty:
        detalle = pd.DataFrame(columns=[
            "cliente", "periodo", "operacion", "fecha", "tipo_comprobante",
            "punto_venta", "numero", "cuit", "razon_social", "neto_gravado",
            "alicuota", "iva", "no_gravado_exento", "percepciones", "total",
            "moneda", "tipo_cambio", "neto_gravado_origen", "iva_origen",
            "no_gravado_exento_origen", "percepciones_origen", "total_origen",
        ])
        resumen_cliente = pd.DataFrame(columns=["cliente", "operacion", "comprobantes", "neto_gravado", "iva", "total"])
        resumen_alicuota = pd.DataFrame(columns=["operacion", "alicuota", "neto_gravado", "iva"])
        return detalle, resumen_cliente, resumen_alicuota

    detalle = comprobantes.copy()
    if "detalle_alicuotas" in detalle:
        varias = detalle["detalle_alicuotas"].map(lambda partidas: isinstance(partidas, (list, tuple)) and len(partidas) > 1)
        # pandas 3 rejects assigning text into a float column, so make the
        # display column capable of holding both a rate and the label.
        detalle["alicuota"] = detalle["alicuota"].astype(object)
        detalle.loc[varias, "alicuota"] = "Varias"
    # NC reducen débito/crédito; ND y facturas suman. Se conserva el importe original en el detalle.
    detalle["factor"] = detalle["tipo_comprobante"].str.startswith("NC").map({True: -1, False: 1})
    detalle["neto_impacto"] = detalle["neto_gravado"] * detalle["factor"]
    detalle["iva_impacto"] = detalle["iva"] * detalle["factor"]
    detalle["total_impacto"] = detalle["total"] * detalle["factor"]
    resumen_cliente = detalle.groupby(["cliente", "operacion"], as_index=False).agg(
        comprobantes=("numero", "count"), neto_gravado=("neto_impacto", "sum"),
        iva=("iva_impacto", "sum"), total=("total_impacto", "sum"),
    )
    filas_alicuota = []
    for _, comprobante in detalle.iterrows():
        partidas = comprobante.get("detalle_alicuotas")
        if not isinstance(partidas, (list, tuple)) or not partidas:
            partidas = [{
                "alicuota": comprobante["alicuota"],
                "neto_gravado": comprobante["neto_gravado"],
                "iva": comprobante["iva"],
            }]
        for partida in partidas:
            filas_alicuota.append({
                "operacion": comprobante["operacion"],
                "alicuota": float(partida["alicuota"]),
                "neto_gravado": float(partida["neto_gravado"]) * comprobante["factor"],
                "iva": float(partida["iva"]) * comprobante["factor"],
            })
    resumen_alicuota = pd.DataFrame(filas_alicuota).groupby(
        ["operacion", "alicuota"], as_index=False
    ).agg(neto_gravado=("neto_gravado", "sum"), iva=("iva", "sum"))
    if "detalle_alicuotas" in detalle:
        detalle["detalle_alicuotas"] = detalle["detalle_alicuotas"].map(
            lambda partidas: json.dumps(partidas, ensure_ascii=False) if isinstance(partidas, (list, tuple)) else "[]"
        )
    detalle = detalle.drop(columns=["factor", "neto_impacto", "iva_impacto", "total_impacto"])
    return detalle, resumen_cliente, resumen_alicuota


def construir_detalle_alicuotas(comprobantes):
    """Desglosa cada comprobante en filas por tasa de IVA para revisión."""
    columnas = [
        "cliente", "periodo", "operacion", "fecha", "tipo_comprobante",
        "punto_venta", "numero", "cuit", "alicuota", "neto_gravado",
        "iva", "moneda", "tipo_cambio",
    ]
    filas = []
    for comprobante in comprobantes.to_dict("records"):
        partidas = comprobante.get("detalle_alicuotas")
        if not isinstance(partidas, (list, tuple)) or not partidas:
            partidas = [{
                "alicuota": comprobante.get("alicuota", 0),
                "neto_gravado": comprobante["neto_gravado"],
                "iva": comprobante["iva"],
            }]
        factor = -1 if str(comprobante["tipo_comprobante"]).startswith("NC") else 1
        for partida in partidas:
            filas.append({
                "cliente": comprobante["cliente"],
                "periodo": comprobante["periodo"],
                "operacion": comprobante["operacion"],
                "fecha": comprobante["fecha"],
                "tipo_comprobante": comprobante["tipo_comprobante"],
                "punto_venta": comprobante["punto_venta"],
                "numero": comprobante["numero"],
                "cuit": comprobante["cuit"],
                "alicuota": partida["alicuota"],
                "neto_gravado": float(partida["neto_gravado"]) * factor,
                "iva": float(partida["iva"]) * factor,
                "moneda": comprobante.get("moneda", "ARS"),
                "tipo_cambio": comprobante.get("tipo_cambio", 1),
            })
    return pd.DataFrame(filas, columns=columnas)


def aplicar_formato_excel(ruta_archivo):
    libro = load_workbook(ruta_archivo)
    azul, celeste, blanco = "1F4E78", "D9EAF7", "FFFFFF"
    for hoja in libro.worksheets:
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        hoja.row_dimensions[1].height = 24
        for celda in hoja[1]:
            celda.fill = PatternFill("solid", fgColor=azul)
            celda.font = Font(bold=True, color=blanco)
            celda.alignment = Alignment(horizontal="center")
        for columna in hoja.columns:
            letra = get_column_letter(columna[0].column)
            hoja.column_dimensions[letra].width = min(max(max(len(str(c.value or "")) for c in columna) + 2, 12), 42)
            encabezado = str(hoja.cell(1, columna[0].column).value or "").casefold()
            if any(p in encabezado for p in ("neto", "iva", "total", "percepcion")):
                for celda in columna[1:]:
                    celda.number_format = '"$"#,##0.00'
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if celda.row % 2 == 0:
                    celda.fill = PatternFill("solid", fgColor=celeste)
        if hoja.max_row > 1:
            hoja.auto_filter.ref = f"A1:{get_column_letter(hoja.max_column)}{hoja.max_row}"
    libro.save(ruta_archivo)
    libro.close()
