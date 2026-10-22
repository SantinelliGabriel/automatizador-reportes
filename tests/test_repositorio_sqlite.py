"""Persistencia versionada y operaciones separadas de confirmación e informe."""

import hashlib
import sqlite3
import unittest
from contextlib import closing
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from automatizador.database import VERSION_ESQUEMA, crear_tabla
from automatizador.procesamiento import confirmar_importacion, generar_informe, previsualizar, procesar
from automatizador.repositorio import RepositorioSQLite
from automatizador.validaciones import validar_comprobantes
from tests.helpers import comprobante, directorio_temporal


class TestRepositorioSQLite(unittest.TestCase):
    def test_ventas_luego_compras_del_mismo_cliente_periodo_reemplaza_ventas(self):
        with directorio_temporal() as carpeta:
            repo = RepositorioSQLite(Path(carpeta) / "comprobantes.db")
            for operacion, numero in (("ventas", 1), ("compras", 2)):
                aceptados, rechazados = validar_comprobantes(pd.DataFrame([
                    comprobante(numero=numero, operacion=operacion),
                ]), "2026-09")
                repo.reemplazar_periodos([{
                    "cliente": "Cliente Sintético", "periodo": "2026-09",
                    "comprobantes": aceptados, "rechazados": rechazados,
                    "archivo_origen": f"{operacion}.csv",
                }])

            actuales = repo.consultar_comprobantes("2026-09")
            self.assertEqual(actuales[["operacion", "numero"]].values.tolist(), [["compras", 2]])
            self.assertEqual(repo.consultar_importaciones()["archivo_origen"].tolist(), [
                "compras.csv", "ventas.csv",
            ])

    def test_dos_archivos_del_mismo_cliente_periodo_en_un_lote_dejan_solo_el_ultimo(self):
        with directorio_temporal() as carpeta:
            repo = RepositorioSQLite(Path(carpeta) / "comprobantes.db")
            lotes = []
            for operacion, numero in (("ventas", 1), ("compras", 2)):
                aceptados, rechazados = validar_comprobantes(pd.DataFrame([
                    comprobante(numero=numero, operacion=operacion),
                ]), "2026-09")
                lotes.append({
                    "cliente": "Cliente Sintético", "periodo": "2026-09",
                    "comprobantes": aceptados, "rechazados": rechazados,
                    "archivo_origen": f"{operacion}.csv",
                })

            ids = repo.reemplazar_periodos(lotes)

            self.assertEqual(len(ids), 2)
            actuales = repo.consultar_comprobantes("2026-09")
            self.assertEqual(actuales[["operacion", "numero"]].values.tolist(), [["compras", 2]])
            self.assertEqual(len(repo.consultar_importaciones()), 2)

    def test_excel_compuesto_y_regenerado_tienen_mismas_hojas_y_celdas(self):
        with directorio_temporal() as carpeta:
            carpeta = Path(carpeta)
            entrada = carpeta / "Cliente Sintético.csv"
            pd.DataFrame([
                comprobante(numero=1),
                comprobante(numero=2, cuit="CUIT inválido"),
            ]).to_csv(entrada, index=False)
            base = carpeta / "comprobantes.db"
            anterior = carpeta / "anterior.xlsx"
            nuevo = carpeta / "regenerado.xlsx"
            procesar(entrada, anterior, base, "2026-09")
            generar_informe("2026-09", nuevo, RepositorioSQLite(base))

            with closing(load_workbook(anterior, read_only=True, data_only=True)) as libro_anterior:
                with closing(load_workbook(nuevo, read_only=True, data_only=True)) as libro_nuevo:
                    self.assertEqual(libro_anterior.sheetnames, libro_nuevo.sheetnames)
                    for nombre in libro_anterior.sheetnames:
                        self.assertEqual(
                            list(libro_anterior[nombre].values), list(libro_nuevo[nombre].values), nombre,
                        )

    def test_confirma_observaciones_y_sha_sin_generar_excel_y_lo_regenera_sin_original(self):
        with directorio_temporal() as carpeta:
            carpeta = Path(carpeta)
            entrada = carpeta / "Cliente Sintético.csv"
            pd.DataFrame([
                comprobante(numero=1),
                comprobante(numero=2, cuit="CUIT inválido"),
            ]).to_csv(entrada, index=False)
            sha256 = hashlib.sha256(entrada.read_bytes()).hexdigest()
            repo = RepositorioSQLite(carpeta / "datos" / "comprobantes.db")
            salida = carpeta / "resultados" / "reporte.xlsx"

            plan = previsualizar(entrada, "2026-09")
            ids = confirmar_importacion(plan, repo)

            self.assertEqual(len(ids), 1)
            self.assertFalse(salida.exists())
            self.assertEqual(len(repo.consultar_comprobantes("2026-09")), 1)
            historial = repo.consultar_importaciones()
            self.assertEqual(historial.iloc[0]["archivo_origen"], entrada.name)
            self.assertEqual(historial.iloc[0]["sha256"], sha256)
            self.assertEqual(historial.iloc[0]["aceptados"], 1)
            self.assertEqual(historial.iloc[0]["rechazados"], 1)
            observaciones = repo.consultar_observaciones(importacion_id=ids[0])
            self.assertEqual(len(observaciones), 1)
            self.assertEqual(observaciones.iloc[0]["fila_excel"], "3")
            self.assertIn("CUIT", observaciones.iloc[0]["motivo"])

            entrada.unlink()
            self.assertEqual(generar_informe("2026-09", salida, repo), salida)
            self.assertTrue(salida.is_file())
            self.assertEqual(len(pd.read_excel(salida, sheet_name="Comprobantes")), 1)
            self.assertEqual(len(pd.read_excel(salida, sheet_name="Observaciones")), 1)

    def test_archivo_totalmente_observado_conserva_comprobantes_previos(self):
        with directorio_temporal() as carpeta:
            carpeta = Path(carpeta)
            entrada = carpeta / "Cliente Sintético.csv"
            repo = RepositorioSQLite(carpeta / "comprobantes.db")
            pd.DataFrame([comprobante(numero=1)]).to_csv(entrada, index=False)
            primer_id = confirmar_importacion(previsualizar(entrada, "2026-09"), repo)[0]

            pd.DataFrame([comprobante(numero=2, cuit="CUIT inválido")]).to_csv(entrada, index=False)
            nuevo_plan = previsualizar(entrada, "2026-09")
            self.assertFalse(nuevo_plan.archivos[0].aplicar)
            segundo_id = confirmar_importacion(nuevo_plan, repo)[0]

            self.assertEqual(repo.consultar_comprobantes("2026-09")["numero"].tolist(), [1])
            self.assertEqual(len(repo.consultar_importaciones()), 2)
            self.assertEqual(len(repo.consultar_observaciones(importacion_id=segundo_id)), 1)
            self.assertEqual(len(repo.consultar_observaciones("2026-09")), 1)
            self.assertTrue(repo.consultar_observaciones(importacion_id=primer_id).empty)

    def test_reimportacion_conserva_observaciones_historicas_y_muestra_las_vigentes(self):
        with directorio_temporal() as carpeta:
            carpeta = Path(carpeta)
            entrada = carpeta / "Cliente Sintético.csv"
            repo = RepositorioSQLite(carpeta / "comprobantes.db")
            pd.DataFrame([
                comprobante(numero=1), comprobante(numero=2, cuit="CUIT inválido"),
            ]).to_csv(entrada, index=False)
            primer_id = confirmar_importacion(previsualizar(entrada, "2026-09"), repo)[0]

            pd.DataFrame([comprobante(numero=3)]).to_csv(entrada, index=False)
            confirmar_importacion(previsualizar(entrada, "2026-09"), repo)

            self.assertEqual(repo.consultar_comprobantes("2026-09")["numero"].tolist(), [3])
            self.assertEqual(len(repo.consultar_observaciones(importacion_id=primer_id)), 1)
            self.assertTrue(repo.consultar_observaciones("2026-09").empty)

    def test_error_en_lote_revierte_tambien_observaciones(self):
        with directorio_temporal() as carpeta:
            repo = RepositorioSQLite(Path(carpeta) / "comprobantes.db")
            buenos, observados = validar_comprobantes(pd.DataFrame([
                comprobante(numero=1), comprobante(numero=2, cuit="CUIT inválido"),
            ]), "2026-09")
            repetidos, _ = validar_comprobantes(pd.DataFrame([comprobante(numero=3)]), "2026-09")
            with self.assertRaises(sqlite3.IntegrityError):
                repo.reemplazar_periodos([
                    {
                        "cliente": "Cliente Sintético Uno", "periodo": "2026-09",
                        "comprobantes": buenos, "rechazados": observados, "archivo_origen": "uno.csv",
                    },
                    {
                        "cliente": "Cliente Sintético Dos", "periodo": "2026-09",
                        "comprobantes": pd.concat([repetidos, repetidos], ignore_index=True),
                        "rechazados": pd.DataFrame(), "archivo_origen": "dos.csv",
                    },
                ])

            self.assertTrue(repo.consultar_importaciones().empty)
            self.assertTrue(repo.consultar_observaciones("2026-09").empty)
            self.assertTrue(repo.consultar_comprobantes("2026-09").empty)

    def test_migra_base_preexistente_sin_perder_comprobantes(self):
        with directorio_temporal() as carpeta:
            base = Path(carpeta) / "comprobantes.db"
            with closing(sqlite3.connect(base)) as conexion:
                conexion.executescript("""CREATE TABLE importaciones (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    fecha_importacion TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    archivo_origen TEXT NOT NULL,
                    cantidad_ventas INTEGER NOT NULL DEFAULT 0,
                    total_centavos INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE comprobantes (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    importacion_id INTEGER NOT NULL REFERENCES importaciones(id),
                    cliente TEXT NOT NULL, periodo TEXT NOT NULL, operacion TEXT NOT NULL,
                    fecha TEXT NOT NULL, tipo_comprobante TEXT NOT NULL,
                    punto_venta INTEGER NOT NULL, numero INTEGER NOT NULL,
                    cuit TEXT NOT NULL, razon_social TEXT NOT NULL,
                    neto_gravado_centavos INTEGER NOT NULL, alicuota REAL NOT NULL,
                    iva_centavos INTEGER NOT NULL, no_gravado_exento_centavos INTEGER NOT NULL,
                    percepciones_centavos INTEGER NOT NULL, total_centavos INTEGER NOT NULL,
                    UNIQUE(cliente, periodo, operacion, tipo_comprobante, punto_venta, numero, cuit)
                );""")
                conexion.execute(
                    "INSERT INTO importaciones(archivo_origen, cantidad_ventas, total_centavos) VALUES (?, ?, ?)",
                    ("sintetico.csv", 1, 12100),
                )
                conexion.execute("""INSERT INTO comprobantes (
                    importacion_id, cliente, periodo, operacion, fecha, tipo_comprobante, punto_venta,
                    numero, cuit, razon_social, neto_gravado_centavos, alicuota, iva_centavos,
                    no_gravado_exento_centavos, percepciones_centavos, total_centavos
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", (
                    1, "Cliente Sintético", "2026-09", "ventas", "2026-09-03", "A", 1,
                    101, "CUIT-SINTETICO", "Empresa Sintética", 10000, 21, 2100, 0, 0, 12100,
                ))
                conexion.commit()

            crear_tabla(base)
            crear_tabla(base)
            with closing(sqlite3.connect(base)) as conexion:
                self.assertEqual(conexion.execute("PRAGMA user_version").fetchone()[0], VERSION_ESQUEMA)
                self.assertEqual(conexion.execute("SELECT COUNT(*) FROM comprobantes").fetchone()[0], 1)
                self.assertEqual(conexion.execute(
                    "SELECT neto_gravado_centavos, iva_centavos, total_centavos, total_origen_centavos "
                    "FROM comprobantes"
                ).fetchone(), (10000, 2100, 12100, 12100))
                self.assertEqual(conexion.execute("SELECT sha256 FROM importaciones").fetchone()[0], "")
                self.assertIsNotNone(conexion.execute(
                    "SELECT name FROM sqlite_master WHERE name = 'observaciones'"
                ).fetchone())

    def test_rechaza_version_futura_sin_alterar_base(self):
        with directorio_temporal() as carpeta:
            base = Path(carpeta) / "comprobantes.db"
            with closing(sqlite3.connect(base)) as conexion:
                conexion.execute("PRAGMA user_version = 99")
            with self.assertRaisesRegex(ValueError, "versión de esquema no admitida"):
                crear_tabla(base)
            with closing(sqlite3.connect(base)) as conexion:
                self.assertEqual(conexion.execute("PRAGMA user_version").fetchone()[0], 99)
                self.assertIsNone(conexion.execute(
                    "SELECT name FROM sqlite_master WHERE name = 'comprobantes'"
                ).fetchone())
