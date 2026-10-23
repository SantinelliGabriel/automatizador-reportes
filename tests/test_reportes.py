import unittest
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from automatizador.database import crear_tabla, consultar_comprobantes, consultar_totales, reemplazar_periodo
from automatizador.procesamiento import generar_informe
from automatizador.reportes import aplicar_formato_excel, construir_hojas
from automatizador.repositorio import RepositorioSQLite
from automatizador.validaciones import validar_comprobantes
from tests.helpers import comprobante, directorio_temporal


class TestReportes(unittest.TestCase):
    def test_formula_llega_como_texto_a_cada_hoja(self):
        with directorio_temporal() as carpeta:
            ruta = Path(carpeta) / "informe.xlsx"
            libro = Workbook()
            for nombre in ("Comprobantes", "Detalle IVA", "Resumen por cliente",
                           "IVA por alícuota", "Observaciones"):
                hoja = libro.active if nombre == "Comprobantes" else libro.create_sheet()
                hoja.title = nombre
                hoja.append(["texto"])
                for valor in ("=1+1", "+SUM(1)", "-SUM(1)", "@SUM(1)"):
                    hoja.append([valor])
            libro.save(ruta)
            libro.close()
            aplicar_formato_excel(ruta)
            libro = load_workbook(ruta, data_only=False)
            for hoja in libro:
                self.assertEqual([(hoja.cell(fila, 1).value, hoja.cell(fila, 1).data_type)
                                  for fila in range(2, 6)],
                                 [(valor, "s") for valor in ("=1+1", "+SUM(1)", "-SUM(1)", "@SUM(1)")])
            libro.close()

    def test_texto_peligroso_se_escribe_literal_en_todas_las_hojas(self):
        filas = [comprobante(numero=1, razon_social="=1+1"),
                 comprobante(numero=2, razon_social="@SUM(1)")]
        aceptados, observados = validar_comprobantes(pd.DataFrame(filas), "2026-09")
        self.assertTrue(observados.empty)
        with directorio_temporal() as carpeta:
            base = Path(carpeta) / "comprobantes.db"
            salida = Path(carpeta) / "reporte.xlsx"
            repo = RepositorioSQLite(base)
            reemplazar_periodo("Cliente Sintético", "2026-09", aceptados, observados, "sintetico.csv", base)
            generar_informe("2026-09", salida, repo)
            libro = load_workbook(salida, data_only=False)
            hoja = libro["Comprobantes"]
            columna = [celda.value for celda in hoja[1]].index("razon_social") + 1
            self.assertEqual([(hoja.cell(fila, columna).value, hoja.cell(fila, columna).data_type)
                              for fila in (2, 3)], [("=1+1", "s"), ("@SUM(1)", "s")])
            libro.close()

    def test_totales_sql_coinciden_con_excel_con_nc_y_nd(self):
        filas = [
            comprobante(numero=1, neto_gravado="100.00", iva="21.00", total="121.00"),
            comprobante(numero=2, tipo_comprobante="NC A", neto_gravado="20.00", iva="4.20", total="24.20"),
            comprobante(numero=3, tipo_comprobante="ND A", neto_gravado="10.00", iva="2.10", total="12.10"),
        ]
        buenos, malos = validar_comprobantes(pd.DataFrame(filas), "2026-09")
        self.assertTrue(malos.empty)
        with directorio_temporal() as carpeta:
            base = Path(carpeta) / "comprobantes.db"
            salida = Path(carpeta) / "reporte.xlsx"
            repo = RepositorioSQLite(base)
            reemplazar_periodo("Cliente Sintético", "2026-09", buenos, malos, "sintetico.csv", base)
            generar_informe("2026-09", salida, repo)
            sql = consultar_totales(base, "2026-09")
            resumen = pd.read_excel(salida, sheet_name="Resumen por cliente")
            alicuotas = pd.read_excel(salida, sheet_name="IVA por alícuota")

        self.assertEqual(sql, {
            "neto_gravado": Decimal("90.00"), "iva": Decimal("18.90"), "total": Decimal("108.90"),
        })
        for columna, total in sql.items():
            self.assertEqual(sum(Decimal(str(x)) for x in resumen[columna]), total)
        for columna in ("neto_gravado", "iva"):
            self.assertEqual(sum(Decimal(str(x)) for x in alicuotas[columna]), sql[columna])

    def test_notas_de_credito_y_debito_ajustan_resumen_y_alicuota(self):
        venta = comprobante()
        nota_credito = comprobante(tipo_comprobante="NC A", numero=102, neto_gravado=10000, iva=2100, total=12100)
        nota_debito = comprobante(
            tipo_comprobante="ND A", numero=103, neto_gravado=5000,
            alicuota=10.5, iva=525, total=5525,
        )
        buenos, malos = validar_comprobantes(
            pd.DataFrame([venta, nota_credito, nota_debito]), "2026-09"
        )
        with directorio_temporal() as carpeta:
            base = Path(carpeta) / "comprobantes.db"
            crear_tabla(base)
            reemplazar_periodo("Cliente Demo", "2026-09", buenos, malos, "demo.xlsx", base)
            detalle, por_cliente, por_alicuota = construir_hojas(consultar_comprobantes(base))
        self.assertEqual(len(detalle), 3)
        self.assertTrue(malos.empty)
        self.assertEqual(por_cliente.iloc[0]["neto_gravado"], 95000)
        self.assertEqual(por_cliente.iloc[0]["iva"], 19425)
        self.assertEqual(por_cliente.iloc[0]["total"], 114425)
        por_tasa = por_alicuota.set_index("alicuota")
        self.assertEqual(por_tasa.loc[21.0, "neto_gravado"], 90000)
        self.assertEqual(por_tasa.loc[21.0, "iva"], 18900)
        self.assertEqual(por_tasa.loc[10.5, "neto_gravado"], 5000)
        self.assertEqual(por_tasa.loc[10.5, "iva"], 525)
