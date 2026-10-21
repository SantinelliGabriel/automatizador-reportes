import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from automatizador.conciliacion import conciliar_banco, exportar_conciliacion
from tests.helpers import comprobante, directorio_temporal


class TestConciliacion(unittest.TestCase):
    def _conciliar(self, movimientos, documentos, descripcion="Descripcion"):
        temporal = directorio_temporal()
        self.addCleanup(temporal.cleanup)
        ruta = Path(temporal.name) / "extracto.csv"
        pd.DataFrame(movimientos).to_csv(ruta, index=False)
        return conciliar_banco(
            ruta,
            pd.DataFrame(documentos),
            "Fecha",
            "Importe",
            descripcion,
        )

    def test_concilia_ventas_compras_y_notas_de_credito_por_signo(self):
        documentos = [
            comprobante(numero=101, total=121),
            comprobante(tipo_comprobante="NC A", numero=102, total=50),
            comprobante(operacion="compras", numero=103, total=100),
        ]
        movimientos = [
            {"Fecha": "03/09/2026", "Importe": 121, "Descripcion": "Factura 1-101"},
            {"Fecha": "03/09/2026", "Importe": -50, "Descripcion": "NC 1-102"},
            {"Fecha": "03/09/2026", "Importe": -100, "Descripcion": "Compra 1-103"},
        ]

        banco, docs = self._conciliar(movimientos, documentos)

        self.assertEqual(banco["Estado"].tolist(), ["Conciliado"] * 3)
        self.assertEqual(docs["Estado"].tolist(), ["Conciliado"] * 3)
        self.assertIn("A 00001-00000101", banco.loc[0, "Comprobante"])

    def test_deja_empates_como_ambiguos_y_movimientos_sin_match(self):
        documentos = [comprobante(numero=n, total=121) for n in (101, 102)]
        movimientos = [
            {"Fecha": "03/09/2026", "Importe": 121, "Descripcion": "Transferencia"},
            {"Fecha": "03/09/2026", "Importe": 999, "Descripcion": "Otro movimiento"},
        ]

        banco, docs = self._conciliar(movimientos, documentos)

        self.assertEqual(banco.loc[0, "Estado"], "Coincidencia ambigua")
        self.assertIn("00000101", banco.loc[0, "Candidatos"])
        self.assertIn("00000102", banco.loc[0, "Candidatos"])
        self.assertEqual(banco.loc[1, "Estado"], "Sin coincidencia")
        self.assertTrue((docs["Estado"] == "Coincidencia ambigua").all())

    def test_sugiere_pago_parcial_con_referencia_a_factura(self):
        documento = comprobante(numero=201, total=121000)
        movimiento = {"Fecha": "03/09/2026", "Importe": 50000, "Descripcion": "Pago factura 1-201"}

        banco, docs = self._conciliar([movimiento], [documento])

        self.assertEqual(banco.loc[0, "Estado"], "Posible pago parcial; revisar")
        self.assertEqual(docs.loc[0, "Estado"], "Referencia detectada; revisar diferencia")

    def test_sugiere_pago_agrupado_sin_conciliarlo_automaticamente(self):
        documentos = [
            comprobante(numero=301, neto_gravado=40000, iva=0, total=40000),
            comprobante(numero=302, neto_gravado=60000, iva=0, total=60000),
        ]
        movimiento = {"Fecha": "03/09/2026", "Importe": 100000, "Descripcion": "Pago agrupado"}

        banco, docs = self._conciliar([movimiento], documentos)

        self.assertEqual(banco.loc[0, "Estado"], "Posible pago agrupado; revisar")
        self.assertIn("00000301", banco.loc[0, "Candidatos"])
        self.assertIn("00000302", banco.loc[0, "Candidatos"])
        self.assertTrue((docs["Estado"] == "Posible pago agrupado; revisar").all())

    def test_sugiere_pago_fraccionado_con_referencias_repetidas(self):
        documento = comprobante(numero=401, total=121000)
        movimientos = [
            {"Fecha": "03/09/2026", "Importe": 50000, "Descripcion": "Pago factura 1-401 parte 1"},
            {"Fecha": "04/09/2026", "Importe": 71000, "Descripcion": "Pago factura 1-401 parte 2"},
        ]

        banco, docs = self._conciliar(movimientos, [documento])

        self.assertTrue((banco["Estado"] == "Posible pago fraccionado; revisar").all())
        self.assertEqual(docs.loc[0, "Estado"], "Posible pago fraccionado; revisar")

    def test_valida_columnas_y_exporta_hojas_de_revision(self):
        documento = comprobante()
        with directorio_temporal() as carpeta:
            ruta = Path(carpeta) / "extracto.csv"
            pd.DataFrame([{"Otra": "dato"}]).to_csv(ruta, index=False)
            with self.assertRaisesRegex(ValueError, "El extracto no tiene las columnas"):
                conciliar_banco(ruta, pd.DataFrame([documento]), "Fecha", "Importe")

            banco = pd.DataFrame([{"Estado": "Sin coincidencia", "Importe": 10}])
            docs = pd.DataFrame([{"Estado": "Sin coincidencia", "numero": 1}])
            salida = Path(carpeta) / "conciliacion.xlsx"
            exportar_conciliacion(salida, banco, docs)
            libro = load_workbook(salida, read_only=True)
            try:
                self.assertEqual(
                    libro.sheetnames,
                    ["Extracto bancario", "Comprobantes", "Pendientes banco", "Pendientes comprobantes"],
                )
            finally:
                libro.close()


if __name__ == "__main__":
    unittest.main()
